import streamlit as st
import pandas as pd
import re
import openpyxl
from openpyxl.styles import PatternFill
import html
import io
from collections import defaultdict
from html.parser import HTMLParser
from spellchecker import SpellChecker

# ---------------------------------------------------------
# 1. Page Configuration
# ---------------------------------------------------------
st.set_page_config(
    page_title="Vizard Guns Data Auditor Pro",
    page_icon="✨",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Initialize SpellChecker
spell = SpellChecker()

# ---------------------------------------------------------
# 2. Apple-Inspired Glassmorphic CSS (Greyish-White + Dark Orange & Green)
# ---------------------------------------------------------
st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700;800&display=swap');

    html, body, [class*="css"] {
        font-family: 'Plus Jakarta Sans', -apple-system, BlinkMacSystemFont, sans-serif;
    }

    /* Overall Background: Smooth Greyish-White Gradient */
    .stApp {
        background: linear-gradient(135deg, #F1F5F9 0%, #E2E8F0 50%, #CBD5E1 100%);
        color: #0F172A;
    }

    /* Main Container Padding */
    .main .block-container {
        padding-top: 2rem;
        padding-bottom: 4rem;
        max-width: 1280px;
    }

    /* Header Title Design */
    .hero-title-container {
        text-align: center;
        padding: 30px 20px 20px 20px;
        margin-bottom: 25px;
        background: rgba(255, 255, 255, 0.7);
        backdrop-filter: blur(20px);
        -webkit-backdrop-filter: blur(20px);
        border-radius: 24px;
        border: 1px solid rgba(255, 255, 255, 0.8);
        box-shadow: 0 20px 40px -15px rgba(0, 0, 0, 0.05);
    }

    .hero-title {
        font-size: 2.75rem;
        font-weight: 800;
        background: linear-gradient(135deg, #EA580C 0%, #D97706 50%, #059669 100%);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        letter-spacing: -0.03em;
        margin-bottom: 8px;
    }

    .hero-subtitle {
        font-size: 1.05rem;
        color: #475569;
        font-weight: 500;
    }

    /* Cards & Containers: Rounded Aesthetic */
    div[data-testid="stForm"], 
    div.stFileUploader, 
    div[data-testid="stVerticalBlock"] > div {
        border-radius: 20px !important;
    }

    .stCard {
        background: rgba(255, 255, 255, 0.75);
        backdrop-filter: blur(16px);
        border: 1px solid rgba(255, 255, 255, 0.9);
        border-radius: 20px;
        padding: 24px;
        box-shadow: 0 10px 30px -10px rgba(0, 0, 0, 0.04);
        margin-bottom: 20px;
    }

    /* Metric Cards Styling: Dark Orange Accent */
    div[data-testid="stMetric"] {
        background: rgba(255, 255, 255, 0.85) !important;
        border: 1px solid rgba(226, 232, 240, 0.8) !important;
        border-radius: 18px !important;
        padding: 18px 22px !important;
        box-shadow: 0 10px 25px -5px rgba(0, 0, 0, 0.03) !important;
        border-top: 4px solid #EA580C !important;
        transition: transform 0.2s ease, box-shadow 0.2s ease;
    }

    div[data-testid="stMetric"]:hover {
        transform: translateY(-2px);
        box-shadow: 0 15px 30px -5px rgba(234, 88, 12, 0.1) !important;
    }

    div[data-testid="stMetricValue"] {
        color: #0F172A !important;
        font-weight: 800 !important;
        font-size: 1.85rem !important;
    }

    div[data-testid="stMetricLabel"] {
        color: #64748B !important;
        font-weight: 600 !important;
        font-size: 0.85rem !important;
        text-transform: uppercase;
        letter-spacing: 0.05em;
    }

    /* Primary Buttons: Dark Orange Gradient */
    .stButton > button[kind="primary"] {
        width: 100%;
        background: linear-gradient(135deg, #EA580C 0%, #C2410C 100%) !important;
        color: #FFFFFF !important;
        font-weight: 700 !important;
        font-size: 1.05rem !important;
        border-radius: 14px !important;
        border: none !important;
        padding: 14px 28px !important;
        box-shadow: 0 8px 20px -4px rgba(234, 88, 12, 0.4) !important;
        transition: all 0.25s ease !important;
    }

    .stButton > button[kind="primary"]:hover {
        transform: translateY(-2px) !important;
        box-shadow: 0 12px 25px -4px rgba(234, 88, 12, 0.5) !important;
        background: linear-gradient(135deg, #F97316 0%, #EA580C 100%) !important;
    }

    /* Download Buttons: Emerald Green Gradient */
    .stDownloadButton > button {
        width: 100%;
        background: linear-gradient(135deg, #059669 0%, #047857 100%) !important;
        color: #FFFFFF !important;
        font-weight: 700 !important;
        font-size: 1.05rem !important;
        border-radius: 14px !important;
        border: none !important;
        padding: 14px 28px !important;
        box-shadow: 0 8px 20px -4px rgba(5, 150, 105, 0.4) !important;
        transition: all 0.25s ease !important;
    }

    .stDownloadButton > button:hover {
        transform: translateY(-2px) !important;
        box-shadow: 0 12px 25px -4px rgba(5, 150, 105, 0.5) !important;
        background: linear-gradient(135deg, #10B981 0%, #059669 100%) !important;
    }

    /* Inputs, Selectboxes, Textareas Aesthetics */
    .stSelectbox > div > div, .stMultiSelect > div > div, .stTextInput > div > div, .stTextArea > div > div {
        border-radius: 12px !important;
        border: 1px solid #CBD5E1 !important;
        background-color: #FFFFFF !important;
        color: #0F172A !important;
    }

    /* Section Headings */
    h2, h3 {
        color: #0F172A !important;
        font-weight: 700 !important;
        letter-spacing: -0.02em;
    }

    /* Progress Bar */
    .stProgress > div > div > div > div {
        background: linear-gradient(90deg, #EA580C 0%, #10B981 100%) !important;
        border-radius: 10px !important;
    }

    /* Custom Footer */
    .footer-text {
        text-align: center;
        color: #94A3B8;
        font-size: 0.85rem;
        margin-top: 40px;
    }
    </style>
""", unsafe_allow_html=True)

# ---------------------------------------------------------
# 3. Header Title Block
# ---------------------------------------------------------
st.markdown("""
    <div class="hero-title-container">
        <div class="hero-title">✨ Vizard Guns Data Auditor </div>
       
    </div>
""", unsafe_allow_html=True)

# ---------------------------------------------------------
# 4. HTML Parser Class for Structure Validation
# ---------------------------------------------------------
class TagBalanceChecker(HTMLParser):
    VOID_TAGS = {'meta', 'img', 'br', 'hr', 'input', 'link', 'area', 'base', 'col', 'embed', 'source', 'track', 'wbr'}

    def __init__(self):
        super().__init__()
        self.stack = []
        self.errors = []

    def handle_starttag(self, tag, attrs):
        if tag.lower() not in self.VOID_TAGS:
            self.stack.append((tag.lower(), self.getpos()))

    def handle_endtag(self, tag):
        tag_lower = tag.lower()
        if tag_lower in self.VOID_TAGS:
            return
            
        if not self.stack:
            self.errors.append(f"Unexpected closing tag </{tag_lower}>")
            return

        expected_tag, _ = self.stack[-1]
        if tag_lower == expected_tag:
            self.stack.pop()
        else:
            stack_tags = [t[0] for t in self.stack]
            if tag_lower in stack_tags:
                while self.stack and self.stack[-1][0] != tag_lower:
                    unclosed_tag, _ = self.stack.pop()
                    self.errors.append(f"Unclosed tag <{unclosed_tag}>")
                if self.stack:
                    self.stack.pop()
            else:
                self.errors.append(f"Mismatched closing tag </{tag_lower}> (expected </{expected_tag}>)")

    def get_audit_summary(self):
        while self.stack:
            unclosed_tag, _ = self.stack.pop()
            self.errors.append(f"Unclosed tag <{unclosed_tag}> at end")
        return self.errors

def validate_html_structure(html_string):
    if not html_string or not isinstance(html_string, str) or not html_string.strip():
        return "Empty Description"
    
    checker = TagBalanceChecker()
    try:
        checker.feed(html_string)
        errors = checker.get_audit_summary()
        if errors:
            return "HTML Tag Errors: " + " | ".join(errors)
        return "HTML Structure Valid"
    except Exception as e:
        return f"HTML Syntax Exception: {str(e)}"

# ---------------------------------------------------------
# 5. Helper Functions for New Checks
# ---------------------------------------------------------
CORRUPTED_PATTERNS = [
    r'â€–', r'â€™', r'â€œ', r'â€', r'Ã©', r'Ã', r'â€\x9d', r'Â',
    r'“', r'”', r'‘', r'’', r'–', r'—', r'×', r'\xa0'
]
CORRUPTED_REGEX = re.compile('|'.join(CORRUPTED_PATTERNS))

def check_corrupted_characters(value):
    if not isinstance(value, str):
        return None
    matches = set(CORRUPTED_REGEX.findall(value))
    if matches:
        return ", ".join([f"'{m}'" for m in matches])
    return None

def clean_word_for_spellcheck(word):
    cleaned = re.sub(r'[^a-zA-Z]', '', word)
    return cleaned

def perform_spell_check(text, whitelist_set):
    if not isinstance(text, str) or not text.strip():
        return []
    
    # Strip HTML tags before tokenizing words
    clean_text = re.sub(r'<[^>]+>', ' ', text)
    tokens = re.findall(r'\b[A-Za-z]{3,}\b', clean_text)
    
    misspelled_words = set()
    for token in tokens:
        # Ignore alphanumeric, upper-case acronyms (e.g., MPN, UPC, SKU)
        if token.isupper() or token.lower() in whitelist_set:
            continue
        
        # Check against pyspellchecker
        unknown = spell.unknown([token.lower()])
        if unknown:
            misspelled_words.add(token)
            
    return list(misspelled_words)

# Default firearms/e-commerce domain terms to whitelist
DEFAULT_WHITELIST = """
glock sig sauer picatinny cerakote mlok luger beretta ruger smith wesson holosun magpul
leupold vortex caliber carbine handguard receiver optics ammo munition cheek riser
choke weaver ambidextrous suppressor threaded lanyard QD baseplate fde
"""

# ---------------------------------------------------------
# 6. File Upload & Mapping Sidebar / Main Layout
# ---------------------------------------------------------
uploaded_file = st.file_uploader("📂 Upload Excel Data Sheet (.xlsx)", type=["xlsx"], key="vizard_standalone_uploader")

if uploaded_file is not None:
    df = pd.read_excel(uploaded_file)
    st.success(f"✓ File successfully loaded ({len(df)} rows found)")
    
    column_options = df.columns.tolist()
    
    # Advanced Settings: Custom Jargon Whitelist for Spell Check
    with st.expander("⚙️ Advanced Audit Settings & Spellcheck Whitelist", expanded=False):
        st.markdown("**Domain Whitelist (Words to ignore during spell check):**")
        custom_whitelist_input = st.text_area(
            "Add custom brand names, SKUs, or industry jargon separated by spaces or newlines:",
            value=DEFAULT_WHITELIST.strip(),
            height=100
        )
        whitelist_words = set(re.findall(r'\b\w+\b', custom_whitelist_input.lower()))

    st.markdown("### 🏷️ Column Mapping Interface")
    
    col1, col2, col3 = st.columns(3)
    with col1:
        title_default = next((i for i, c in enumerate(column_options) if 'title' in c.lower()), 0)
        title_col = st.selectbox("Select Title Column:", column_options, index=title_default, key="vizard_title_select")
    with col2:
        desc_default = next((i for i, c in enumerate(column_options) if 'desc' in c.lower()), 0)
        desc_col = st.selectbox("Select Description Column:", column_options, index=desc_default, key="vizard_desc_select")
    with col3:
        brand_default = next((i for i, c in enumerate(column_options) if 'brand' in c.lower() or 'mfg' in c.lower() or 'manufacturer' in c.lower()), 0)
        brand_col = st.selectbox("Select Brand Column (Optional):", ["None"] + column_options,
                                 index=(column_options.index(column_options[brand_default]) + 1) if ('brand' in column_options[brand_default].lower() or 'mfg' in column_options[brand_default].lower() or 'manufacturer' in column_options[brand_default].lower()) else 0, key="vizard_brand_select")

    col4, col5 = st.columns(2)
    with col4:
        upc_default = next((i for i, c in enumerate(column_options) if 'upc' in c.lower()), 0)
        upc_col = st.selectbox("Select UPC Column (Optional):", ["None"] + column_options, 
                              index=(column_options.index(column_options[upc_default]) + 1) if 'upc' in column_options[upc_default].lower() else 0, key="vizard_upc_select")
    with col5:
        mpn_default = next((i for i, c in enumerate(column_options) if 'mpn' in c.lower() or 'model' in c.lower()), 0)
        mpn_col = st.selectbox("Select MPN/Model Column (Optional):", ["None"] + column_options, 
                              index=(column_options.index(column_options[mpn_default]) + 1) if ('mpn' in column_options[mpn_default].lower() or 'model' in column_options[mpn_default].lower()) else 0, key="vizard_mpn_select")

    st.markdown("<br>", unsafe_allow_html=True)

    if st.button("🚀 Run Data Audit", key="vizard_run_btn", type="primary"):
        desc_position = df.columns.get_loc(desc_col)
        
        audit_cols = [
            'Extracted H2 Title', 'Title Verification Comment', 
            'Title Dash Audit', 'Brand Prefix Verification',
            'Extracted Desc UPC', 'UPC Verification Comment',
            'Extracted Desc MPN', 'MPN Verification Comment',
            'Title Duplicate Status', 'HTML Tag Audit Comment',
            'Corrupted Characters Detected', 'File-Wide Spelling Errors'
        ]
        
        for c in audit_cols:
            if c in df.columns:
                df = df.drop(columns=[c])
                
        for idx, col_name in enumerate(audit_cols, 1):
            df.insert(desc_position + idx, col_name, None)

        excel_buffer = io.BytesIO()
        df.to_excel(excel_buffer, index=False)
        excel_buffer.seek(0)
        
        wb = openpyxl.load_workbook(excel_buffer)
        ws = wb.active

        headers = [cell.value for cell in ws[1]]
        title_col_idx = headers.index(title_col) + 1
        desc_col_idx = headers.index(desc_col) + 1
        brand_col_idx = headers.index(brand_col) + 1 if brand_col != "None" else None
        upc_col_idx = headers.index(upc_col) + 1 if upc_col != "None" else None
        mpn_col_idx = headers.index(mpn_col) + 1 if mpn_col != "None" else None

        ext_h2_idx = headers.index('Extracted H2 Title') + 1
        title_comment_idx = headers.index('Title Verification Comment') + 1
        title_dash_idx = headers.index('Title Dash Audit') + 1
        brand_prefix_idx = headers.index('Brand Prefix Verification') + 1
        ext_upc_idx = headers.index('Extracted Desc UPC') + 1
        upc_comment_idx = headers.index('UPC Verification Comment') + 1
        ext_mpn_idx = headers.index('Extracted Desc MPN') + 1
        mpn_comment_idx = headers.index('MPN Verification Comment') + 1
        dup_comment_idx = headers.index('Title Duplicate Status') + 1
        html_comment_idx = headers.index('HTML Tag Audit Comment') + 1
        corrupt_char_idx = headers.index('Corrupted Characters Detected') + 1
        spelling_err_idx = headers.index('File-Wide Spelling Errors') + 1

        # Color Fills
        red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")     # Soft Red
        yellow_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")  # Soft Orange/Yellow
        orange_fill = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")  # Dark Orange Soft Accent

        progress_bar = st.progress(0.0)
        status_text = st.empty()

        total_rows = ws.max_row
        data_rows = max(1, total_rows - 1)

        # Build Duplicate Title Map
        title_tracker = defaultdict(list)
        for r_idx in range(2, total_rows + 1):
            raw_title_val = str(ws.cell(row=r_idx, column=title_col_idx).value or "").strip()
            if raw_title_val:
                title_tracker[raw_title_val.lower()].append(r_idx)

        mismatch_h2_count = 0
        dash_error_count = 0
        brand_error_count = 0
        mismatch_upc_count = 0
        mismatch_mpn_count = 0
        duplicate_title_rows = 0
        html_error_rows = 0
        corrupted_char_rows = 0
        spelling_error_rows = 0

        # Non-standard dash regex (contains dashes like – or —, but not standard ASCII hyphen -)
        NON_STANDARD_DASH_REGEX = re.compile(r'[\u2010-\u2015\u2212\u2013\u2014]')

        for row_idx in range(2, total_rows + 1):
            status_text.text(f"Auditing Row {row_idx-1} of {total_rows-1}...")
            
            raw_title = str(ws.cell(row=row_idx, column=title_col_idx).value or "").strip()
            desc_text = str(ws.cell(row=row_idx, column=desc_col_idx).value or "")
            brand_text = str(ws.cell(row=row_idx, column=brand_col_idx).value or "").strip() if brand_col_idx else ""
            orig_upc = str(ws.cell(row=row_idx, column=upc_col_idx).value or "").strip().replace('.0', '') if upc_col_idx else ""
            orig_mpn = str(ws.cell(row=row_idx, column=mpn_col_idx).value or "").strip().replace('.0', '') if mpn_col_idx else ""

            # Check 1: H2 Title Verification
            h2_match = re.search(r'<h2[^>]*>(.*?)<\/h2>', desc_text, re.IGNORECASE | re.DOTALL)
            extracted_h2 = html.unescape(re.sub(r'<[^>]+>', '', h2_match.group(1))).strip() if h2_match else ""
            ws.cell(row=row_idx, column=ext_h2_idx).value = extracted_h2

            if not extracted_h2:
                mismatch_h2_count += 1
                ws.cell(row=row_idx, column=title_comment_idx).value = "Missing H2 Tag in Description"
                ws.cell(row=row_idx, column=ext_h2_idx).fill = red_fill
            else:
                clean_title = re.sub(r'\s+', ' ', raw_title).strip()
                clean_h2 = re.sub(r'\s+', ' ', extracted_h2).strip()

                if clean_title.lower() == clean_h2.lower():
                    ws.cell(row=row_idx, column=title_comment_idx).value = "H2 Title Matched"
                else:
                    mismatch_h2_count += 1
                    ws.cell(row=row_idx, column=title_comment_idx).value = f"Mismatch: Title ({clean_title}) vs H2 ({clean_h2})"
                    ws.cell(row=row_idx, column=ext_h2_idx).fill = red_fill

            # Check 2: Standardize Dash Check (Title Column)
            if NON_STANDARD_DASH_REGEX.search(raw_title):
                dash_error_count += 1
                ws.cell(row=row_idx, column=title_dash_idx).value = "Non-Standard Dash Detected (Replace with Hyphen '-')"
                ws.cell(row=row_idx, column=title_dash_idx).fill = orange_fill
                ws.cell(row=row_idx, column=title_col_idx).fill = orange_fill
            else:
                ws.cell(row=row_idx, column=title_dash_idx).value = "Valid (No Non-Standard Dashes)"

            # Check 3: Brand Prefix Verification
            if brand_col_idx and brand_text:
                if not raw_title.lower().startswith(brand_text.lower()):
                    brand_error_count += 1
                    ws.cell(row=row_idx, column=brand_prefix_idx).value = f"Title does not start with Brand '{brand_text}'"
                    ws.cell(row=row_idx, column=brand_prefix_idx).fill = yellow_fill
                else:
                    ws.cell(row=row_idx, column=brand_prefix_idx).value = "Brand Prefix Validated"

            # Check 4: UPC & MPN Verification
            upc_match = re.search(r'UPC:\s*<\/strong>\s*([0-9a-zA-Z]+)', desc_text, re.IGNORECASE) or \
                        re.search(r'UPC:\s*([0-9a-zA-Z]+)', desc_text, re.IGNORECASE)
            extracted_upc = upc_match.group(1).strip() if upc_match else ""
            ws.cell(row=row_idx, column=ext_upc_idx).value = extracted_upc
            if extracted_upc:
                ws.cell(row=row_idx, column=ext_upc_idx).number_format = '@'

            if upc_col_idx and orig_upc:
                if extracted_upc and orig_upc != extracted_upc:
                    mismatch_upc_count += 1
                    ws.cell(row=row_idx, column=upc_comment_idx).value = f"UPC Mismatch: Sheet ({orig_upc}) vs Desc ({extracted_upc})"
                    ws.cell(row=row_idx, column=ext_upc_idx).fill = red_fill
                elif not extracted_upc:
                    ws.cell(row=row_idx, column=upc_comment_idx).value = "UPC Missing in Description HTML"
                else:
                    ws.cell(row=row_idx, column=upc_comment_idx).value = "UPC Validated"
            elif extracted_upc:
                ws.cell(row=row_idx, column=upc_comment_idx).value = "UPC Found in Description"

            mpn_match = re.search(r'(?:MPN|Model):\s*<\/strong>\s*([0-9a-zA-Z-]+)', desc_text, re.IGNORECASE) or \
                        re.search(r'(?:MPN|Model):\s*([0-9a-zA-Z-]+)', desc_text, re.IGNORECASE) or \
                        re.search(r'<strong>MPN<\/strong>\s*([0-9a-zA-Z-]+)', desc_text, re.IGNORECASE)
            extracted_mpn = mpn_match.group(1).strip() if mpn_match else ""
            ws.cell(row=row_idx, column=ext_mpn_idx).value = extracted_mpn
            if extracted_mpn:
                ws.cell(row=row_idx, column=ext_mpn_idx).number_format = '@'

            if mpn_col_idx and orig_mpn:
                if extracted_mpn and orig_mpn.lower() != extracted_mpn.lower():
                    mismatch_mpn_count += 1
                    ws.cell(row=row_idx, column=mpn_comment_idx).value = f"MPN Mismatch: Sheet ({orig_mpn}) vs Desc ({extracted_mpn})"
                    ws.cell(row=row_idx, column=ext_mpn_idx).fill = red_fill
                elif not extracted_mpn:
                    ws.cell(row=row_idx, column=mpn_comment_idx).value = "MPN Missing in Description HTML"
                else:
                    ws.cell(row=row_idx, column=mpn_comment_idx).value = "MPN Validated"
            elif extracted_mpn:
                ws.cell(row=row_idx, column=mpn_comment_idx).value = "MPN Found in Description"

            # Check 5: Title Duplication
            all_matches = title_tracker.get(raw_title.lower(), [])
            if len(all_matches) > 1:
                duplicate_title_rows += 1
                rows_str = ", ".join(map(str, all_matches))
                ws.cell(row=row_idx, column=dup_comment_idx).value = f"DUPLICATE TITLE (Found in {len(all_matches)} rows: {rows_str})"
                ws.cell(row=row_idx, column=dup_comment_idx).fill = yellow_fill
                ws.cell(row=row_idx, column=title_col_idx).fill = yellow_fill
            else:
                ws.cell(row=row_idx, column=dup_comment_idx).value = "Unique Title"

            # Check 6: HTML Structure Check
            html_audit_res = validate_html_structure(desc_text)
            ws.cell(row=row_idx, column=html_comment_idx).value = html_audit_res
            if "HTML Tag Errors" in html_audit_res or "Syntax Exception" in html_audit_res:
                html_error_rows += 1
                ws.cell(row=row_idx, column=html_comment_idx).fill = red_fill
                ws.cell(row=row_idx, column=desc_col_idx).fill = red_fill

            # Check 7: File-Wide Corrupted Character Scan
            row_corruptions = []
            row_spelling_errs = []

            for c_idx in range(1, len(headers) + 1):
                col_header = str(headers[c_idx - 1])
                # Skip audit generated columns
                if col_header in audit_cols:
                    continue

                cell_val = ws.cell(row=row_idx, column=c_idx).value
                if cell_val:
                    # Corrupted Chars
                    corrupt_found = check_corrupted_characters(str(cell_val))
                    if corrupt_found:
                        row_corruptions.append(f"[{col_header}]: Contains {corrupt_found}")
                        ws.cell(row=row_idx, column=c_idx).fill = orange_fill

                    # Spell Check
                    spelling_found = perform_spell_check(str(cell_val), whitelist_words)
                    if spelling_found:
                        formatted_words = ", ".join([f'"{w}"' for w in spelling_found])
                        row_spelling_errs.append(f"[{col_header}]: {formatted_words}")

            if row_corruptions:
                corrupted_char_rows += 1
                ws.cell(row=row_idx, column=corrupt_char_idx).value = " | ".join(row_corruptions)
                ws.cell(row=row_idx, column=corrupt_char_idx).fill = orange_fill
            else:
                ws.cell(row=row_idx, column=corrupt_char_idx).value = "Clean"

            if row_spelling_errs:
                spelling_error_rows += 1
                ws.cell(row=row_idx, column=spelling_err_idx).value = " | ".join(row_spelling_errs)
                ws.cell(row=row_idx, column=spelling_err_idx).fill = yellow_fill
            else:
                ws.cell(row=row_idx, column=spelling_err_idx).value = "Clean"

            progress_val = float(min(1.0, max(0.0, (row_idx - 1) / data_rows)))
            progress_bar.progress(progress_val)
            
        status_text.text("🎉 Audit Completed Successfully!")

        # ---------------------------------------------------------
        # 7. Dashboard Overview Section
        # ---------------------------------------------------------
        st.markdown("<br><h2 style='text-align: center;'>📊 Executive Audit Overview</h2>", unsafe_allow_html=True)
        
        total_errors = (mismatch_h2_count + dash_error_count + brand_error_count + 
                        mismatch_upc_count + mismatch_mpn_count + duplicate_title_rows + 
                        html_error_rows + corrupted_char_rows + spelling_error_rows)
        
        m1, m2, m3 = st.columns(3)
        m1.metric("Total Rows Scanned", data_rows)
        m2.metric("Total Issue Flags", total_errors)
        m3.metric("Data Quality Score", f"{max(0, round(((data_rows - total_errors) / max(1, data_rows)) * 100, 1))}%")

        st.markdown("<br>", unsafe_allow_html=True)
        
        dash_col1, dash_col2 = st.columns([1, 1])

        error_data = pd.DataFrame({
            "Area of Audit": [
                "H2 Title Mismatches / Missing",
                "Non-Standard Title Dashes",
                "Brand Prefix Violations",
                "UPC Code Mismatches",
                "MPN / Model Code Mismatches",
                "Duplicate Product Titles",
                "Broken HTML Structure",
                "Corrupted Character Encoding",
                "File-Wide Spelling Errors"
            ],
            "Issue Count": [
                mismatch_h2_count,
                dash_error_count,
                brand_error_count,
                mismatch_upc_count,
                mismatch_mpn_count,
                duplicate_title_rows,
                html_error_rows,
                corrupted_char_rows,
                spelling_error_rows
            ]
        })

        with dash_col1:
            st.markdown("### 📋 Error Breakdown by Area")
            st.dataframe(
                error_data, 
                column_config={
                    "Area of Audit": "Area of Audit",
                    "Issue Count": st.column_config.NumberColumn("Flagged Count")
                },
                use_container_width=True,
                hide_index=True
            )

        with dash_col2:
            st.markdown("### 📈 Visual Issue Distribution")
            st.bar_chart(
                data=error_data.set_index("Area of Audit"),
                y="Issue Count",
                color="#EA580C",
                use_container_width=True
            )

        st.markdown("<br>", unsafe_allow_html=True)

        out_buffer = io.BytesIO()
        wb.save(out_buffer)
        out_buffer.seek(0)
        
        st.download_button(
            label="📥 Download Vizard Guns Audited File",
            data=out_buffer.getvalue(),
            file_name="Vizard_Guns_Checked_Data.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="vizard_download_btn",
            type="primary"
        )

# Footer
st.markdown('<div class="footer-text">Vizard Guns Data Auditor • Designed for High-Precision E-Commerce Verification</div>', unsafe_allow_html=True)
